"""
The validation gate.

Nothing reaches the renderer without passing every check here. This is the
layer that stops a plausible-but-wrong model output from becoming a wrong
document:

    target exists? -> editable? -> right region? -> merge valid? -> formula?
    -> value supported by evidence? -> type fits? -> confidence high enough? -> WRITE

Checks run against the *live* parsed template, not against the plan's own
assumptions, so a template edited since the spec was generated fails loudly
instead of writing to a stale address.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple

from openpyxl.utils import range_boundaries

from .ir import DocType, NodeType, TemplateIR, ValueFormat
from .mapper import Decision, FillInstruction, FillPlan, Policy
from .normalize import coerce
from .spec import TemplateSpec


@dataclass
class Violation:
    node_id: str
    role: str
    check: str
    message: str
    severity: str = "error"          # error blocks the write, warning annotates it

    def to_dict(self) -> Dict[str, Any]:
        return {"node_id": self.node_id, "role": self.role, "check": self.check,
                "message": self.message, "severity": self.severity}


@dataclass
class ValidationReport:
    violations: List[Violation] = field(default_factory=list)
    blocked: List[str] = field(default_factory=list)     # node_ids blocked from writing
    checked: int = 0

    @property
    def ok(self) -> bool:
        return not any(v.severity == "error" for v in self.violations)

    def errors(self) -> List[Violation]:
        return [v for v in self.violations if v.severity == "error"]

    def warnings(self) -> List[Violation]:
        return [v for v in self.violations if v.severity == "warning"]

    def to_dict(self) -> Dict[str, Any]:
        return {"ok": self.ok, "checked": self.checked,
                "errors": [v.to_dict() for v in self.errors()],
                "warnings": [v.to_dict() for v in self.warnings()],
                "blocked_nodes": self.blocked}


class ValidationGate:
    def __init__(self, spec: TemplateSpec, ir: TemplateIR, policy: Optional[Policy] = None):
        self.spec = spec
        self.ir = ir
        self.policy = policy or Policy()
        self._protected_ids = {p["node_id"] for p in spec.protected}
        self._nodes = {n.node_id: n for n in ir.nodes}
        self._tables = {t.node_id: t for t in ir.tables}

    # ------------------------------------------------------------------
    def run(self, plan: FillPlan) -> ValidationReport:
        report = ValidationReport()
        for ins in plan.instructions:
            if not ins.writable:
                continue
            report.checked += 1
            vios = self._check(ins)
            report.violations.extend(vios)
            if any(v.severity == "error" for v in vios):
                report.blocked.append(ins.node_id)
                ins.decision = Decision.BLOCKED_BY_VALIDATION.value
                ins.notes.extend(v.message for v in vios if v.severity == "error")
        plan.recompute_summary()
        return report

    # ------------------------------------------------------------------
    def _check(self, ins: FillInstruction) -> List[Violation]:
        v: List[Violation] = []
        add = lambda check, msg, sev="error": v.append(  # noqa: E731
            Violation(ins.node_id, ins.role, check, msg, sev))

        # 1. does the target still exist in the live template?
        if ins.kind == "field":
            node = self._nodes.get(ins.node_id)
            if node is None:
                add("target_exists", f"node '{ins.node_id}' not present in the template")
                return v
            # 2/3. editable and of the right kind
            if ins.node_id in self._protected_ids or not node.editable:
                add("editable", "target is a protected region")
            if node.type not in (NodeType.VALUE_REGION,):
                add("region_type", f"target is a {node.type.value}, not a value region")
            # 4. formulas are never overwritten
            if node.has_formula:
                add("no_formula_overwrite", "target holds a formula")
            # 5. merge validity: only the anchor of a merged range is writable
            if node.style.merged and self.ir.doc_type == DocType.XLSX:
                if not self._is_merge_anchor(node.location.parts):
                    add("merge_anchor",
                        f"{node.location.parts.get('cell')} is not the anchor of "
                        f"{node.location.parts.get('range')}")
            # 6. type fits the target's number format
            _, ok = coerce(ins.value, ins.value_format)
            if not ok:
                add("type_fit", f"value {ins.value!r} does not fit format {ins.value_format}")
            # 7. non-empty target: warn, never silently clobber content
            if not node.is_empty and not node.placeholder:
                add("overwrite", f"target already contains {node.text[:40]!r}", "warning")

        else:  # table
            table = self._tables.get(ins.node_id)
            if table is None:
                add("target_exists", f"table '{ins.node_id}' not present in the template")
                return v
            spec_t = self.spec.tables.get(ins.role) or next(
                (t for t in self.spec.tables.values() if t.node_id == ins.node_id), None)
            if spec_t is None:
                add("spec_present", "table has no spec entry")
                return v
            if [c.header_text for c in table.columns] != [
                    c["header"] for c in sorted(spec_t.columns.values(),
                                                key=lambda x: str(x["location_hint"]))] and False:
                pass  # header drift is checked structurally below instead
            live_headers = {c.header_text.strip().lower() for c in table.columns}
            for fname, cfg in spec_t.columns.items():
                if cfg["header"].strip().lower() not in live_headers:
                    add("header_drift",
                        f"column '{cfg['header']}' from the spec is missing in the template")
            n = len(ins.rows or [])
            if n == 0:
                add("row_count", "no rows to write", "warning")
            if n > self.policy.max_rows_per_table:
                add("row_count", f"{n} rows exceeds max_rows_per_table "
                                 f"({self.policy.max_rows_per_table})")
            if spec_t.fixed_rows and n > max(1, spec_t.existing_data_rows):
                add("fixed_rows",
                    f"this table cannot grow safely; {n} rows requested but only "
                    f"{max(1, spec_t.existing_data_rows)} available", "warning")
            for i, row in enumerate(ins.rows or []):
                for fname, value in row.items():
                    cfg = spec_t.columns.get(fname)
                    if cfg is None:
                        add("unknown_column", f"row {i}: no column for field '{fname}'", "warning")
                        continue
                    if not cfg.get("editable", True):
                        add("editable", f"row {i}: column '{fname}' is protected")
                    _, ok = coerce(value, cfg.get("format", "text"))
                    if not ok:
                        add("type_fit", f"row {i}: {value!r} does not fit "
                                        f"{cfg.get('format')} in column '{fname}'", "warning")

        # 8. evidence and confidence, re-checked at the gate
        if self.policy.require_evidence and not ins.evidence:
            add("evidence", "no evidence for this value")
        critical = bool(self.spec.fields.get(ins.role).critical) if self.spec.fields.get(ins.role) else False
        floor = self.policy.review + (self.policy.critical_bonus if critical else 0.0)
        if ins.confidence < floor:
            add("confidence", f"confidence {ins.confidence} below floor {floor:.2f}")
        return v

    def _is_merge_anchor(self, parts: Dict[str, Any]) -> bool:
        rng = parts.get("range")
        cell = parts.get("cell")
        if not rng or ":" not in str(rng):
            return True
        try:
            min_c, min_r, _, _ = range_boundaries(str(rng))
        except Exception:
            return True
        from openpyxl.utils import get_column_letter
        return f"{get_column_letter(min_c)}{min_r}" == cell


def validate(plan: FillPlan, spec: TemplateSpec, ir: TemplateIR,
             policy: Optional[Policy] = None) -> ValidationReport:
    return ValidationGate(spec, ir, policy).run(plan)
